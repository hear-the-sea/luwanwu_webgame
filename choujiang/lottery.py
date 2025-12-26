#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
抽奖程序 - 豪华版 v2
翻牌式抽奖 + 彩带烟花特效 + 全屏模式 + 手动翻牌 + 洗牌效果
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import random
import csv
import json
import os
import math
from typing import List, Tuple, Optional
from dataclasses import dataclass, field

# 尝试导入 openpyxl 用于读取 Excel
try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False


# ==================== 颜色主题 ====================
class Theme:
    """颜色主题配置"""
    BG_DARK = '#0d1b2a'
    BG_MEDIUM = '#1b263b'
    BG_LIGHT = '#415a77'

    GOLD = '#ffd60a'
    GOLD_LIGHT = '#ffee32'
    RED = '#ef233c'
    RED_LIGHT = '#ff6b6b'
    GREEN = '#06d6a0'
    CYAN = '#00b4d8'
    PURPLE = '#9d4edd'

    TEXT_PRIMARY = '#ffffff'
    TEXT_SECONDARY = '#adb5bd'
    TEXT_DARK = '#1b263b'

    CARD_BACK = '#ef233c'
    CARD_FRONT = '#ffd60a'

    FIREWORK_COLORS = ['#ff6b6b', '#ffd60a', '#06d6a0', '#00b4d8', '#9d4edd', '#ff9f1c', '#e76f51']


@dataclass
class Prize:
    """奖项数据类"""
    name: str
    count: int
    winners: List[str] = field(default_factory=list)

    @property
    def remaining(self) -> int:
        return self.count - len(self.winners)


# ==================== 粒子特效系统 ====================
class Particle:
    """单个粒子"""
    def __init__(self, x: float, y: float, vx: float, vy: float,
                 color: str, size: float = 4, lifetime: int = 60,
                 particle_type: str = 'circle'):
        self.x = x
        self.y = y
        self.vx = vx
        self.vy = vy
        self.color = color
        self.size = size
        self.lifetime = lifetime
        self.max_lifetime = lifetime
        self.particle_type = particle_type

    def update(self, gravity: float = 0.15):
        self.x += self.vx
        self.y += self.vy
        self.vy += gravity
        self.vx *= 0.99
        self.lifetime -= 1
        return self.lifetime > 0

    @property
    def alpha(self) -> float:
        return max(0, self.lifetime / self.max_lifetime)


class ParticleSystem:
    """粒子系统"""
    def __init__(self, canvas: tk.Canvas):
        self.canvas = canvas
        self.particles: List[Particle] = []
        self.running = False
        self.particle_ids = []

    def create_firework(self, x: float, y: float):
        num_particles = random.randint(30, 50)
        for _ in range(num_particles):
            angle = random.uniform(0, 2 * math.pi)
            speed = random.uniform(3, 12)
            self.particles.append(Particle(
                x, y, math.cos(angle) * speed, math.sin(angle) * speed,
                random.choice(Theme.FIREWORK_COLORS),
                random.uniform(3, 8), random.randint(40, 80), 'circle'
            ))

    def create_confetti(self, x: float, y: float, count: int = 50):
        for _ in range(count):
            self.particles.append(Particle(
                x + random.uniform(-100, 100), y + random.uniform(-50, 50),
                random.uniform(-3, 3), random.uniform(-8, -2),
                random.choice(Theme.FIREWORK_COLORS),
                random.uniform(8, 15), random.randint(80, 150), 'rect'
            ))

    def create_celebration(self, width: int, height: int):
        for _ in range(5):
            self.create_firework(random.uniform(width*0.2, width*0.8),
                               random.uniform(height*0.2, height*0.5))
        for _ in range(3):
            self.create_confetti(random.uniform(0, width), -20, 30)

    def update_and_draw(self):
        for pid in self.particle_ids:
            self.canvas.delete(pid)
        self.particle_ids.clear()

        alive = []
        for p in self.particles:
            if p.update():
                alive.append(p)
                size = p.size * p.alpha
                if p.particle_type == 'circle':
                    pid = self.canvas.create_oval(p.x-size, p.y-size, p.x+size, p.y+size,
                                                  fill=p.color, outline='')
                else:
                    pid = self.canvas.create_rectangle(p.x-size/2, p.y-size,
                                                       p.x+size/2, p.y+size,
                                                       fill=p.color, outline='')
                self.particle_ids.append(pid)
        self.particles = alive

        if self.particles:
            self.canvas.after(16, self.update_and_draw)
        else:
            self.running = False

    def start(self):
        if not self.running and self.particles:
            self.running = True
            self.update_and_draw()

    def clear(self):
        self.particles.clear()
        for pid in self.particle_ids:
            self.canvas.delete(pid)
        self.particle_ids.clear()
        self.running = False


# ==================== 可点击翻牌卡片 ====================
class FlipCard(tk.Canvas):
    """可点击的翻牌卡片"""

    def __init__(self, master, width=220, height=140, name="", on_flip_complete=None, **kwargs):
        super().__init__(master, width=width, height=height,
                        bg=Theme.BG_DARK, highlightthickness=0, **kwargs)
        self.width = width
        self.height = height
        self.name = name
        self.is_flipped = False
        self.animation_running = False
        self.flip_progress = 0
        self.on_flip_complete = on_flip_complete
        self.particle_system = ParticleSystem(self)

        # 绑定点击事件
        self.bind('<Button-1>', self._on_click)
        self.bind('<Enter>', self._on_enter)
        self.bind('<Leave>', self._on_leave)

        self.draw_card_back()

    def _on_click(self, event):
        """点击翻牌"""
        if not self.is_flipped and not self.animation_running:
            self.animate_flip()

    def _on_enter(self, event):
        """鼠标进入 - 高亮效果"""
        if not self.is_flipped:
            self.config(cursor='hand2')
            self._draw_hover_effect()

    def _on_leave(self, event):
        """鼠标离开"""
        if not self.is_flipped:
            self.config(cursor='')
            self.draw_card_back()

    def _draw_hover_effect(self):
        """绘制悬停效果"""
        self.delete("all")
        cx, cy = self.width // 2, self.height // 2
        pad = 6

        # 发光边框
        self.create_rectangle(pad-2, pad-2, self.width-pad+2, self.height-pad+2,
                             fill='', outline=Theme.GOLD_LIGHT, width=3)
        self.create_rectangle(pad, pad, self.width-pad, self.height-pad,
                             fill=Theme.CARD_BACK, outline=Theme.RED_LIGHT, width=3)

        inner_pad = 18
        self.create_rectangle(inner_pad, inner_pad, self.width-inner_pad, self.height-inner_pad,
                             fill='', outline=Theme.RED_LIGHT, width=1, dash=(5, 3))

        self.create_text(cx+2, cy+2, text="?", font=('Arial Black', 48, 'bold'), fill='#000')
        self.create_text(cx, cy, text="?", font=('Arial Black', 48, 'bold'), fill=Theme.TEXT_PRIMARY)

        # 点击提示
        self.create_text(cx, self.height - 18, text="点击翻牌",
                        font=('Microsoft YaHei', 10), fill=Theme.GOLD)

    def draw_card_back(self):
        """绘制卡片背面"""
        self.delete("all")
        cx, cy = self.width // 2, self.height // 2
        pad = 8

        self.create_rectangle(pad+3, pad+3, self.width-pad+3, self.height-pad+3,
                             fill='#000000', outline='')
        self.create_rectangle(pad, pad, self.width-pad, self.height-pad,
                             fill=Theme.CARD_BACK, outline=Theme.RED_LIGHT, width=3)

        inner_pad = 20
        self.create_rectangle(inner_pad, inner_pad, self.width-inner_pad, self.height-inner_pad,
                             fill='', outline=Theme.RED_LIGHT, width=1, dash=(5, 3))

        self.create_text(cx+2, cy+2, text="?", font=('Arial Black', 48, 'bold'), fill='#000')
        self.create_text(cx, cy, text="?", font=('Arial Black', 48, 'bold'), fill=Theme.TEXT_PRIMARY)

        deco_size = 10
        for x, y in [(pad+12, pad+12), (self.width-pad-12, pad+12),
                     (pad+12, self.height-pad-12), (self.width-pad-12, self.height-pad-12)]:
            self.create_oval(x-deco_size/2, y-deco_size/2, x+deco_size/2, y+deco_size/2,
                           fill=Theme.GOLD, outline='')

    def draw_card_front(self):
        """绘制卡片正面"""
        self.delete("all")
        cx, cy = self.width // 2, self.height // 2
        pad = 8

        self.create_rectangle(pad+3, pad+3, self.width-pad+3, self.height-pad+3,
                             fill='#000000', outline='')
        self.create_rectangle(pad, pad, self.width-pad, self.height-pad,
                             fill=Theme.CARD_FRONT, outline=Theme.GOLD_LIGHT, width=3)

        inner_pad = 15
        self.create_rectangle(inner_pad, inner_pad, self.width-inner_pad, self.height-inner_pad,
                             fill='', outline=Theme.TEXT_DARK, width=2)

        font_size = 28 if len(self.name) <= 4 else 22 if len(self.name) <= 6 else 16
        self.create_text(cx+1, cy+1, text=self.name,
                        font=('Microsoft YaHei', font_size, 'bold'), fill='#b8860b')
        self.create_text(cx, cy, text=self.name,
                        font=('Microsoft YaHei', font_size, 'bold'), fill=Theme.TEXT_DARK)

        # 星星装饰
        for x, y in [(pad+15, pad+15), (self.width-pad-15, pad+15),
                     (pad+15, self.height-pad-15), (self.width-pad-15, self.height-pad-15)]:
            self._draw_star(x, y, 7, Theme.RED)

    def _draw_star(self, x: float, y: float, size: float, color: str):
        points = []
        for i in range(5):
            angle = math.radians(i * 72 - 90)
            points.extend([x + size * math.cos(angle), y + size * math.sin(angle)])
            angle = math.radians(i * 72 - 90 + 36)
            points.extend([x + size * 0.4 * math.cos(angle), y + size * 0.4 * math.sin(angle)])
        self.create_polygon(points, fill=color, outline='')

    def animate_flip(self):
        """执行翻牌动画"""
        if self.animation_running:
            return
        self.animation_running = True
        self.flip_progress = 0
        self._do_flip_animation()

    def _do_flip_animation(self):
        self.flip_progress += 1
        total_frames = 20

        if self.flip_progress <= total_frames // 2:
            progress = self.flip_progress / (total_frames // 2)
            self._draw_scaled_back(1 - progress)
        else:
            progress = (self.flip_progress - total_frames // 2) / (total_frames // 2)
            self._draw_scaled_front(progress)

        if self.flip_progress < total_frames:
            self.after(25, self._do_flip_animation)
        else:
            self.draw_card_front()
            self.is_flipped = True
            self.animation_running = False
            self.particle_system.create_firework(self.width // 2, self.height // 2)
            self.particle_system.start()
            if self.on_flip_complete:
                self.on_flip_complete(self)

    def _draw_scaled_back(self, scale: float):
        self.delete("all")
        cx = self.width // 2
        half_w = (self.width - 16) // 2 * max(scale, 0.02)
        self.create_rectangle(cx - half_w, 8, cx + half_w, self.height - 8,
                             fill=Theme.CARD_BACK, outline=Theme.RED_LIGHT, width=2)
        if scale > 0.2:
            self.create_text(cx, self.height//2, text="?",
                            font=('Arial Black', max(10, int(48 * scale)), 'bold'),
                            fill=Theme.TEXT_PRIMARY)

    def _draw_scaled_front(self, scale: float):
        self.delete("all")
        cx = self.width // 2
        half_w = (self.width - 16) // 2 * max(scale, 0.02)
        self.create_rectangle(cx - half_w, 8, cx + half_w, self.height - 8,
                             fill=Theme.CARD_FRONT, outline=Theme.GOLD_LIGHT, width=2)
        if scale > 0.3:
            font_size = 28 if len(self.name) <= 4 else 20
            self.create_text(cx, self.height//2, text=self.name,
                            font=('Microsoft YaHei', max(10, int(font_size * scale)), 'bold'),
                            fill=Theme.TEXT_DARK)

    def reset(self):
        """重置卡片"""
        self.is_flipped = False
        self.animation_running = False
        self.particle_system.clear()
        self.draw_card_back()


# ==================== 洗牌动画卡片 ====================
class ShuffleCard(tk.Canvas):
    """洗牌动画用的卡片"""
    def __init__(self, master, x, y, width=120, height=80, **kwargs):
        super().__init__(master, width=width, height=height,
                        bg=Theme.BG_DARK, highlightthickness=0, **kwargs)
        self.card_width = width
        self.card_height = height
        self.base_x = x
        self.base_y = y
        self.current_x = x
        self.current_y = y
        self.target_x = x
        self.target_y = y
        self._draw()

    def _draw(self):
        self.delete("all")
        pad = 4
        self.create_rectangle(pad, pad, self.card_width-pad, self.card_height-pad,
                             fill=Theme.CARD_BACK, outline=Theme.RED_LIGHT, width=2)
        self.create_text(self.card_width//2, self.card_height//2, text="?",
                        font=('Arial Black', 28, 'bold'), fill=Theme.TEXT_PRIMARY)

    def move_to(self, x, y):
        self.target_x = x
        self.target_y = y

    def update_position(self) -> bool:
        """更新位置，返回是否还在移动"""
        dx = (self.target_x - self.current_x) * 0.2
        dy = (self.target_y - self.current_y) * 0.2

        if abs(dx) < 0.5 and abs(dy) < 0.5:
            self.current_x = self.target_x
            self.current_y = self.target_y
            return False

        self.current_x += dx
        self.current_y += dy
        return True


# ==================== 主应用程序 ====================
class LotteryApp:
    """抽奖程序主应用"""

    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("幸运抽奖系统")
        self.root.geometry("1400x900")
        self.root.configure(bg=Theme.BG_DARK)
        self.root.minsize(1200, 700)

        # 数据
        self.participants: List[str] = []
        self.available: List[str] = []
        self.prizes: List[Prize] = []
        self.current_prize_idx: int = 0
        self.flip_cards: List[FlipCard] = []
        self.pending_winners: List[str] = []  # 已分配到卡片但还未翻开确认的中奖者
        self.cards_ready = False  # 卡片是否已准备好（洗牌后）

        # 全屏
        self.is_fullscreen = False
        self.fullscreen_window = None

        # 洗牌动画
        self.shuffle_cards = []
        self.shuffling = False

        self._create_ui()

        self.root.bind('<F11>', lambda e: self._toggle_fullscreen())
        self.root.bind('<Escape>', lambda e: self._exit_fullscreen())

        self._load_config()

    def _create_ui(self):
        self._create_header()

        self.main_frame = tk.Frame(self.root, bg=Theme.BG_DARK)
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)

        self._create_control_panel()
        self._create_display_area()
        self._create_status_bar()

    def _create_header(self):
        header = tk.Frame(self.root, bg=Theme.BG_MEDIUM, height=70)
        header.pack(fill=tk.X)
        header.pack_propagate(False)

        title_frame = tk.Frame(header, bg=Theme.BG_MEDIUM)
        title_frame.pack(side=tk.LEFT, padx=30)

        tk.Label(title_frame, text="LUCKY DRAW",
                font=('Arial Black', 28, 'bold'),
                fg=Theme.GOLD, bg=Theme.BG_MEDIUM).pack(side=tk.LEFT)

        tk.Label(title_frame, text="  幸运抽奖",
                font=('Microsoft YaHei', 16),
                fg=Theme.TEXT_SECONDARY, bg=Theme.BG_MEDIUM).pack(side=tk.LEFT, pady=(8, 0))

        btn_frame = tk.Frame(header, bg=Theme.BG_MEDIUM)
        btn_frame.pack(side=tk.RIGHT, padx=20)

        tk.Button(btn_frame, text="⛶ 全屏模式",
                 font=('Microsoft YaHei', 11),
                 command=self._toggle_fullscreen,
                 bg=Theme.PURPLE, fg=Theme.TEXT_PRIMARY,
                 activebackground='#b366e0',
                 relief=tk.FLAT, padx=15, pady=5,
                 cursor='hand2').pack(side=tk.LEFT, padx=5)

    def _create_control_panel(self):
        self.control_frame = tk.Frame(self.main_frame, bg=Theme.BG_MEDIUM, width=350)
        self.control_frame.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15))
        self.control_frame.pack_propagate(False)

        inner = tk.Frame(self.control_frame, bg=Theme.BG_MEDIUM)
        inner.pack(fill=tk.BOTH, expand=True, padx=15, pady=15)

        self._create_participant_section(inner)
        self._create_prize_section(inner)
        self._create_winners_section(inner)

    def _create_participant_section(self, parent):
        section = tk.LabelFrame(parent, text=" 📋 参与者名单 ",
                               font=('Microsoft YaHei', 12, 'bold'),
                               fg=Theme.CYAN, bg=Theme.BG_MEDIUM)
        section.pack(fill=tk.X, pady=(0, 12))

        btn_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        btn_frame.pack(fill=tk.X, padx=10, pady=8)

        tk.Button(btn_frame, text="📂 导入名单",
                 font=('Microsoft YaHei', 10),
                 command=self._import_participants,
                 bg=Theme.GREEN, fg=Theme.TEXT_DARK,
                 relief=tk.FLAT, padx=10, pady=4,
                 cursor='hand2').pack(side=tk.LEFT, padx=(0, 8))

        tk.Button(btn_frame, text="🗑️ 清空",
                 font=('Microsoft YaHei', 10),
                 command=self._clear_participants,
                 bg=Theme.RED, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=10, pady=4,
                 cursor='hand2').pack(side=tk.LEFT)

        self.participant_count_label = tk.Label(section,
                                                text="共 0 人 | 可抽 0 人",
                                                font=('Microsoft YaHei', 11, 'bold'),
                                                fg=Theme.GOLD, bg=Theme.BG_MEDIUM)
        self.participant_count_label.pack(pady=6)

        list_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        list_frame.pack(fill=tk.X, padx=10, pady=(0, 8))

        self.participant_listbox = tk.Listbox(list_frame, height=4,
                                             font=('Microsoft YaHei', 10),
                                             bg=Theme.BG_DARK, fg=Theme.TEXT_PRIMARY,
                                             selectbackground=Theme.CYAN,
                                             relief=tk.SOLID, bd=1)
        scrollbar = tk.Scrollbar(list_frame, command=self.participant_listbox.yview)
        self.participant_listbox.config(yscrollcommand=scrollbar.set)
        self.participant_listbox.pack(side=tk.LEFT, fill=tk.X, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def _create_prize_section(self, parent):
        section = tk.LabelFrame(parent, text=" 🏆 奖项设置 ",
                               font=('Microsoft YaHei', 12, 'bold'),
                               fg=Theme.GOLD, bg=Theme.BG_MEDIUM)
        section.pack(fill=tk.X, pady=(0, 12))

        # 输入区域 - 两行布局更清晰
        input_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        input_frame.pack(fill=tk.X, padx=10, pady=8)

        # 第一行：奖项名称
        row1 = tk.Frame(input_frame, bg=Theme.BG_MEDIUM)
        row1.pack(fill=tk.X, pady=2)

        tk.Label(row1, text="奖项名称:",
                font=('Microsoft YaHei', 10),
                fg=Theme.TEXT_PRIMARY, bg=Theme.BG_MEDIUM).pack(side=tk.LEFT)

        self.prize_name_entry = tk.Entry(row1, width=15,
                                        font=('Microsoft YaHei', 11),
                                        bg='#ffffff', fg='#000000',
                                        insertbackground='#000000',
                                        relief=tk.SOLID, bd=1)
        self.prize_name_entry.pack(side=tk.LEFT, padx=5, ipady=3)

        tk.Label(row1, text="人数:",
                font=('Microsoft YaHei', 10),
                fg=Theme.TEXT_PRIMARY, bg=Theme.BG_MEDIUM).pack(side=tk.LEFT, padx=(10, 0))

        self.prize_count_entry = tk.Entry(row1, width=5,
                                         font=('Microsoft YaHei', 11),
                                         bg='#ffffff', fg='#000000',
                                         insertbackground='#000000',
                                         relief=tk.SOLID, bd=1)
        self.prize_count_entry.pack(side=tk.LEFT, padx=5, ipady=3)

        # 第二行：添加按钮
        row2 = tk.Frame(input_frame, bg=Theme.BG_MEDIUM)
        row2.pack(fill=tk.X, pady=(5, 0))

        add_btn = tk.Button(row2, text="✓ 确认添加奖项",
                 font=('Microsoft YaHei', 10, 'bold'),
                 command=self._add_prize,
                 bg=Theme.GREEN, fg=Theme.TEXT_DARK,
                 relief=tk.FLAT, padx=15, pady=3,
                 cursor='hand2')
        add_btn.pack(side=tk.LEFT)

        # 绑定回车键确认
        self.prize_name_entry.bind('<Return>', lambda e: self._add_prize())
        self.prize_count_entry.bind('<Return>', lambda e: self._add_prize())

        # 快速添加
        quick_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        quick_frame.pack(fill=tk.X, padx=10, pady=4)

        tk.Label(quick_frame, text="快捷:",
                font=('Microsoft YaHei', 9),
                fg=Theme.TEXT_SECONDARY, bg=Theme.BG_MEDIUM).pack(side=tk.LEFT)

        for text, name, count in [("特等奖", "特等奖", 1), ("一等奖", "一等奖", 1),
                                  ("二等奖", "二等奖", 3), ("三等奖", "三等奖", 5)]:
            tk.Button(quick_frame, text=text,
                     font=('Microsoft YaHei', 9),
                     command=lambda n=name, c=count: self._quick_add_prize(n, c),
                     bg=Theme.BG_LIGHT, fg=Theme.TEXT_PRIMARY,
                     relief=tk.FLAT, padx=6,
                     cursor='hand2').pack(side=tk.LEFT, padx=2)

        # 奖项列表
        list_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        list_frame.pack(fill=tk.X, padx=10, pady=4)

        self.prize_listbox = tk.Listbox(list_frame, height=4,
                                       font=('Microsoft YaHei', 11),
                                       bg=Theme.BG_DARK, fg=Theme.TEXT_PRIMARY,
                                       selectbackground=Theme.GOLD,
                                       selectforeground=Theme.TEXT_DARK,
                                       relief=tk.SOLID, bd=1,
                                       exportselection=False)
        self.prize_listbox.pack(fill=tk.X)
        self.prize_listbox.bind('<<ListboxSelect>>', self._on_prize_select)

        # 删除按钮
        tk.Button(section, text="🗑️ 删除选中奖项",
                 font=('Microsoft YaHei', 10),
                 command=self._delete_prize,
                 bg=Theme.RED, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=12, pady=4,
                 cursor='hand2').pack(pady=8)

    def _create_winners_section(self, parent):
        section = tk.LabelFrame(parent, text=" 🎁 中奖名单 ",
                               font=('Microsoft YaHei', 12, 'bold'),
                               fg=Theme.RED_LIGHT, bg=Theme.BG_MEDIUM)
        section.pack(fill=tk.BOTH, expand=True)

        text_frame = tk.Frame(section, bg=Theme.BG_MEDIUM)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=8)

        self.winners_text = tk.Text(text_frame,
                                   font=('Microsoft YaHei', 10),
                                   bg=Theme.BG_DARK, fg=Theme.TEXT_PRIMARY,
                                   relief=tk.SOLID, bd=1,
                                   state=tk.DISABLED, wrap=tk.WORD)
        scrollbar = tk.Scrollbar(text_frame, command=self.winners_text.yview)
        self.winners_text.config(yscrollcommand=scrollbar.set)
        self.winners_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        tk.Button(section, text="📤 导出名单",
                 font=('Microsoft YaHei', 10),
                 command=self._export_winners,
                 bg=Theme.GREEN, fg=Theme.TEXT_DARK,
                 relief=tk.FLAT, padx=12,
                 cursor='hand2').pack(pady=8)

    def _create_display_area(self):
        self.display_frame = tk.Frame(self.main_frame, bg=Theme.BG_DARK)
        self.display_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        # 奖项标题
        prize_frame = tk.Frame(self.display_frame, bg=Theme.BG_DARK)
        prize_frame.pack(fill=tk.X, pady=15)

        self.current_prize_label = tk.Label(prize_frame, text="请先设置奖项",
                                           font=('Microsoft YaHei', 38, 'bold'),
                                           fg=Theme.GOLD, bg=Theme.BG_DARK)
        self.current_prize_label.pack()

        self.remaining_label = tk.Label(prize_frame, text="",
                                       font=('Microsoft YaHei', 14),
                                       fg=Theme.TEXT_SECONDARY, bg=Theme.BG_DARK)
        self.remaining_label.pack(pady=(5, 0))

        # 提示标签
        self.hint_label = tk.Label(prize_frame, text="",
                                  font=('Microsoft YaHei', 12),
                                  fg=Theme.CYAN, bg=Theme.BG_DARK)
        self.hint_label.pack(pady=(10, 0))

        # 展示画布
        self.display_canvas = tk.Canvas(self.display_frame, bg=Theme.BG_DARK, highlightthickness=0)
        self.display_canvas.pack(fill=tk.BOTH, expand=True, padx=20)

        self.particle_system = ParticleSystem(self.display_canvas)

        self.cards_container = tk.Frame(self.display_canvas, bg=Theme.BG_DARK)
        self.canvas_window = self.display_canvas.create_window(0, 0, window=self.cards_container, anchor='center')

        self.display_canvas.bind('<Configure>', self._on_canvas_resize)

        # 按钮区
        btn_frame = tk.Frame(self.display_frame, bg=Theme.BG_DARK)
        btn_frame.pack(pady=15)

        self.shuffle_btn = tk.Button(btn_frame, text="🔀 洗牌发牌",
                                    font=('Microsoft YaHei', 14, 'bold'),
                                    command=self._shuffle_animation,
                                    bg=Theme.PURPLE, fg=Theme.TEXT_PRIMARY,
                                    relief=tk.FLAT, padx=25, pady=10,
                                    cursor='hand2')
        self.shuffle_btn.pack(side=tk.LEFT, padx=10)

        self.draw_one_btn = tk.Button(btn_frame, text="🎲 翻一张",
                                     font=('Microsoft YaHei', 14, 'bold'),
                                     command=self._draw_one,
                                     bg=Theme.RED, fg=Theme.TEXT_PRIMARY,
                                     relief=tk.FLAT, padx=25, pady=10,
                                     cursor='hand2')
        self.draw_one_btn.pack(side=tk.LEFT, padx=10)

        self.draw_all_btn = tk.Button(btn_frame, text="🎯 全部翻开",
                                     font=('Microsoft YaHei', 14, 'bold'),
                                     command=self._draw_all,
                                     bg=Theme.GREEN, fg=Theme.TEXT_DARK,
                                     relief=tk.FLAT, padx=25, pady=10,
                                     cursor='hand2')
        self.draw_all_btn.pack(side=tk.LEFT, padx=10)

        # 导航
        nav_frame = tk.Frame(self.display_frame, bg=Theme.BG_DARK)
        nav_frame.pack(pady=10)

        tk.Button(nav_frame, text="◀ 上一奖项",
                 font=('Microsoft YaHei', 11),
                 command=self._prev_prize,
                 bg=Theme.BG_LIGHT, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=15, pady=5,
                 cursor='hand2').pack(side=tk.LEFT, padx=8)

        tk.Button(nav_frame, text="下一奖项 ▶",
                 font=('Microsoft YaHei', 11),
                 command=self._next_prize,
                 bg=Theme.BG_LIGHT, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=15, pady=5,
                 cursor='hand2').pack(side=tk.LEFT, padx=8)

        tk.Button(nav_frame, text="🔄 重置当前",
                 font=('Microsoft YaHei', 11),
                 command=self._reset_current_prize,
                 bg='#ff9500', fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=15, pady=5,
                 cursor='hand2').pack(side=tk.LEFT, padx=8)

    def _on_canvas_resize(self, event):
        self.display_canvas.coords(self.canvas_window, event.width // 2, event.height // 2)

    def _create_status_bar(self):
        status_bar = tk.Frame(self.root, bg=Theme.BG_MEDIUM, height=35)
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)
        status_bar.pack_propagate(False)

        self.status_label = tk.Label(status_bar,
                                    text="就绪 · 按 F11 全屏 · 点击卡片翻牌",
                                    font=('Microsoft YaHei', 10),
                                    fg=Theme.TEXT_SECONDARY, bg=Theme.BG_MEDIUM)
        self.status_label.pack(side=tk.LEFT, padx=20, pady=8)

    # ==================== 洗牌动画 ====================
    def _get_active_canvas(self):
        """获取当前活动的画布（全屏或普通）"""
        if self.is_fullscreen and hasattr(self, 'fs_canvas'):
            return self.fs_canvas
        return self.display_canvas

    def _shuffle_animation(self):
        """执行洗牌动画"""
        if self.shuffling or not self.available:
            return

        self.shuffling = True
        self._clear_cards()
        self._update_hint("🔀 洗牌中...")

        # 在画布上创建洗牌卡片（使用当前活动的画布）
        canvas = self._get_active_canvas()
        w, h = canvas.winfo_width(), canvas.winfo_height()

        num_cards = min(12, len(self.available))
        card_w, card_h = 100, 65

        # 初始位置 - 堆叠在中央
        cx, cy = w // 2, h // 2

        self.shuffle_cards = []
        for i in range(num_cards):
            card = ShuffleCard(canvas, cx - card_w//2, cy - card_h//2, card_w, card_h)
            card.place(x=cx - card_w//2 + i*2, y=cy - card_h//2 + i*2)
            self.shuffle_cards.append(card)

        self._shuffle_step(0)

    def _shuffle_step(self, step: int):
        """洗牌动画步骤"""
        canvas = self._get_active_canvas()
        w, h = canvas.winfo_width(), canvas.winfo_height()
        cx, cy = w // 2, h // 2

        if step < 20:  # 散开
            for i, card in enumerate(self.shuffle_cards):
                angle = (i / len(self.shuffle_cards)) * 2 * math.pi + step * 0.3
                radius = 80 + step * 8
                tx = cx + math.cos(angle) * radius - card.card_width // 2
                ty = cy + math.sin(angle) * radius - card.card_height // 2
                card.move_to(tx, ty)
        elif step < 50:  # 旋转
            for i, card in enumerate(self.shuffle_cards):
                angle = (i / len(self.shuffle_cards)) * 2 * math.pi + step * 0.2
                radius = 150
                tx = cx + math.cos(angle) * radius - card.card_width // 2
                ty = cy + math.sin(angle) * radius - card.card_height // 2
                card.move_to(tx, ty)
        else:  # 收回
            for i, card in enumerate(self.shuffle_cards):
                tx = cx - card.card_width // 2 + random.randint(-30, 30)
                ty = cy - card.card_height // 2 + random.randint(-20, 20)
                card.move_to(tx, ty)

        # 更新位置
        still_moving = False
        for card in self.shuffle_cards:
            if card.update_position():
                still_moving = True
            card.place(x=int(card.current_x), y=int(card.current_y))

        if step < 60:
            self._get_active_window().after(30, lambda: self._shuffle_step(step + 1))
        else:
            # 清理洗牌动画卡片
            for card in self.shuffle_cards:
                card.destroy()
            self.shuffle_cards.clear()
            self.shuffling = False

            # 随机打乱顺序
            random.shuffle(self.available)

            # 展示抽奖卡片
            self._prepare_draw_cards()

    # ==================== 全屏模式 ====================
    def _toggle_fullscreen(self):
        if self.is_fullscreen:
            self._exit_fullscreen()
        else:
            self._enter_fullscreen()

    def _save_card_states(self) -> list:
        """保存所有卡片的状态（名字和是否已翻开）"""
        states = []
        for card in self.flip_cards:
            states.append({
                'name': card.name,
                'flipped': card.is_flipped
            })
        return states

    def _restore_cards_with_states(self, states: list):
        """根据保存的状态恢复卡片"""
        if not states:
            return

        names = [s['name'] for s in states]
        self._create_draw_cards(names)

        # 恢复翻开状态
        for i, state in enumerate(states):
            if state['flipped'] and i < len(self.flip_cards):
                card = self.flip_cards[i]
                card.is_flipped = True
                card.draw_card_front()

        # 更新 pending_winners（只包含未翻开的）
        self.pending_winners = [s['name'] for s in states if not s['flipped']]
        self.cards_ready = True

        # 更新提示
        flipped_count = sum(1 for s in states if s['flipped'])
        remaining = len(states) - flipped_count
        if remaining > 0:
            hint = f"已翻开 {flipped_count} 张，还剩 {remaining} 张待翻开" if flipped_count > 0 else f"👆 共 {len(states)} 张卡片，点击翻牌"
        else:
            hint = "🎊 本轮抽奖完成！可继续洗牌抽取或切换奖项"
        return hint

    def _enter_fullscreen(self):
        if self.is_fullscreen:
            return

        # 保存当前卡片状态（包括已翻开的）
        card_states = self._save_card_states() if self.flip_cards else []

        # 清除普通模式的卡片（不清除 pending_winners，因为我们已经保存了状态）
        self.flip_cards.clear()
        for child in self.cards_container.winfo_children():
            child.destroy()

        self.is_fullscreen = True

        self.fullscreen_window = tk.Toplevel(self.root)
        self.fullscreen_window.attributes('-fullscreen', True)
        self.fullscreen_window.configure(bg=Theme.BG_DARK)
        self.fullscreen_window.bind('<Escape>', lambda e: self._exit_fullscreen())
        self.fullscreen_window.bind('<F11>', lambda e: self._exit_fullscreen())

        self._create_fullscreen_content()

        # 如果之前有卡片，延迟重新创建（等待窗口完全显示）
        if card_states:
            saved_states = card_states.copy()
            def recreate_cards():
                hint = self._restore_cards_with_states(saved_states)
                # 更新位置
                self.fullscreen_window.update_idletasks()
                w, h = self.fs_canvas.winfo_width(), self.fs_canvas.winfo_height()
                self.fs_canvas.coords(self.fs_canvas_window, w // 2, h // 2)
                if hasattr(self, 'fs_hint_label'):
                    self.fs_hint_label.config(text=hint)
                self.hint_label.config(text=hint)
            self.fullscreen_window.after(150, recreate_cards)

    def _create_fullscreen_content(self):
        win = self.fullscreen_window

        self.fs_prize_label = tk.Label(win, text=self._get_prize_text(),
                                       font=('Microsoft YaHei', 64, 'bold'),
                                       fg=Theme.GOLD, bg=Theme.BG_DARK)
        self.fs_prize_label.pack(pady=(50, 15))

        self.fs_remaining_label = tk.Label(win, text=self._get_remaining_text(),
                                          font=('Microsoft YaHei', 22),
                                          fg=Theme.TEXT_SECONDARY, bg=Theme.BG_DARK)
        self.fs_remaining_label.pack()

        self.fs_hint_label = tk.Label(win, text="",
                                     font=('Microsoft YaHei', 16),
                                     fg=Theme.CYAN, bg=Theme.BG_DARK)
        self.fs_hint_label.pack(pady=(15, 0))

        self.fs_canvas = tk.Canvas(win, bg=Theme.BG_DARK, highlightthickness=0)
        self.fs_canvas.pack(fill=tk.BOTH, expand=True, padx=50, pady=20)

        self.fs_particle_system = ParticleSystem(self.fs_canvas)

        self.fs_cards_container = tk.Frame(self.fs_canvas, bg=Theme.BG_DARK)
        self.fs_canvas_window = self.fs_canvas.create_window(0, 0, window=self.fs_cards_container, anchor='center')
        self.fs_canvas.bind('<Configure>', self._on_fs_canvas_resize)

        btn_frame = tk.Frame(win, bg=Theme.BG_DARK)
        btn_frame.pack(pady=25)

        tk.Button(btn_frame, text="🔀 洗牌发牌",
                 font=('Microsoft YaHei', 18, 'bold'),
                 command=self._shuffle_animation,
                 bg=Theme.PURPLE, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=35, pady=12,
                 cursor='hand2').pack(side=tk.LEFT, padx=15)

        tk.Button(btn_frame, text="🎲 翻一张",
                 font=('Microsoft YaHei', 18, 'bold'),
                 command=self._draw_one,
                 bg=Theme.RED, fg=Theme.TEXT_PRIMARY,
                 relief=tk.FLAT, padx=35, pady=12,
                 cursor='hand2').pack(side=tk.LEFT, padx=15)

        tk.Button(btn_frame, text="🎯 全部翻开",
                 font=('Microsoft YaHei', 18, 'bold'),
                 command=self._draw_all,
                 bg=Theme.GREEN, fg=Theme.TEXT_DARK,
                 relief=tk.FLAT, padx=35, pady=12,
                 cursor='hand2').pack(side=tk.LEFT, padx=15)

        nav_frame = tk.Frame(win, bg=Theme.BG_DARK)
        nav_frame.pack(pady=15)

        for text, cmd in [("◀ 上一奖项", self._prev_prize),
                          ("下一奖项 ▶", self._next_prize),
                          ("退出全屏", self._exit_fullscreen)]:
            tk.Button(nav_frame, text=text,
                     font=('Microsoft YaHei', 12),
                     command=cmd,
                     bg=Theme.BG_LIGHT, fg=Theme.TEXT_PRIMARY,
                     relief=tk.FLAT, padx=18, pady=6,
                     cursor='hand2').pack(side=tk.LEFT, padx=8)

    def _on_fs_canvas_resize(self, event):
        self.fs_canvas.coords(self.fs_canvas_window, event.width // 2, event.height // 2)

    def _exit_fullscreen(self):
        if not self.is_fullscreen:
            return

        # 保存当前卡片状态（包括已翻开的）
        card_states = self._save_card_states() if self.flip_cards else []

        # 清除全屏模式的卡片引用（它们会随窗口销毁）
        self.flip_cards.clear()

        self.is_fullscreen = False
        if self.fullscreen_window:
            self.fullscreen_window.destroy()
            self.fullscreen_window = None

        # 如果之前有卡片，延迟重新创建（等待布局更新）
        if card_states:
            saved_states = card_states.copy()
            def recreate_cards():
                hint = self._restore_cards_with_states(saved_states)
                # 更新位置
                self.root.update_idletasks()
                w, h = self.display_canvas.winfo_width(), self.display_canvas.winfo_height()
                self.display_canvas.coords(self.canvas_window, w // 2, h // 2)
                self.hint_label.config(text=hint)
            self.root.after(100, recreate_cards)

    def _get_prize_text(self) -> str:
        if not self.prizes:
            return "请先设置奖项"
        return f"🏆 {self.prizes[self.current_prize_idx].name}"

    def _get_remaining_text(self) -> str:
        if not self.prizes:
            return ""
        p = self.prizes[self.current_prize_idx]
        return f"剩余 {p.remaining} 名 | 已抽 {len(p.winners)} 名"

    def _update_fullscreen(self):
        if self.is_fullscreen and self.fullscreen_window:
            self.fs_prize_label.config(text=self._get_prize_text())
            self.fs_remaining_label.config(text=self._get_remaining_text())

    # ==================== 数据操作 ====================
    def _import_participants(self):
        filetypes = [("所有格式", "*.xlsx *.xls *.csv"), ("Excel", "*.xlsx *.xls"), ("CSV", "*.csv")]
        if not EXCEL_SUPPORT:
            filetypes = [("CSV", "*.csv")]

        filepath = filedialog.askopenfilename(title="选择名单", filetypes=filetypes)
        if not filepath:
            return

        try:
            participants = []
            if filepath.endswith('.csv'):
                with open(filepath, 'r', encoding='utf-8-sig') as f:
                    for row in csv.reader(f):
                        if row and row[0].strip():
                            participants.append(row[0].strip())
            else:
                wb = openpyxl.load_workbook(filepath)
                for row in wb.active.iter_rows(min_row=1, max_col=1, values_only=True):
                    if row[0]:
                        participants.append(str(row[0]).strip())

            headers = ['姓名', '名字', 'Name', 'name', '参与者', '员工']
            if participants and participants[0] in headers:
                participants = participants[1:]
            participants = list(dict.fromkeys(participants))

            self.participants = participants
            self.available = participants.copy()
            self._update_participant_list()
            self.status_label.config(text=f"✅ 导入 {len(participants)} 人")
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def _clear_participants(self):
        if self.participants and messagebox.askyesno("确认", "清空所有参与者？"):
            self.participants = []
            self.available = []
            self._update_participant_list()

    def _update_participant_list(self):
        self.participant_listbox.delete(0, tk.END)
        for p in self.participants[:30]:
            self.participant_listbox.insert(tk.END, p)
        if len(self.participants) > 30:
            self.participant_listbox.insert(tk.END, f"... 共 {len(self.participants)} 人")
        self.participant_count_label.config(
            text=f"共 {len(self.participants)} 人 | 可抽 {len(self.available)} 人"
        )

    def _add_prize(self):
        name = self.prize_name_entry.get().strip()
        count_str = self.prize_count_entry.get().strip()

        if not name:
            messagebox.showwarning("提示", "请输入奖项名称")
            self.prize_name_entry.focus_set()
            return

        try:
            count = int(count_str)
            if count <= 0:
                raise ValueError()
        except:
            messagebox.showwarning("提示", "请输入有效人数（正整数）")
            self.prize_count_entry.focus_set()
            return

        if any(p.name == name for p in self.prizes):
            messagebox.showwarning("提示", f"'{name}' 已存在")
            return

        self.prizes.append(Prize(name=name, count=count))
        self._update_prize_list()

        # 清空输入框
        self.prize_name_entry.delete(0, tk.END)
        self.prize_count_entry.delete(0, tk.END)
        self.prize_name_entry.focus_set()
        self.status_label.config(text=f"已添加奖项: {name} ({count}人)")

    def _quick_add_prize(self, name: str, count: int):
        if any(p.name == name for p in self.prizes):
            messagebox.showwarning("提示", f"'{name}' 已存在")
            return
        self.prizes.append(Prize(name=name, count=count))
        self._update_prize_list()
        self.status_label.config(text=f"已添加: {name}")

    def _delete_prize(self):
        sel = self.prize_listbox.curselection()
        if not sel:
            messagebox.showwarning("提示", "请先选择要删除的奖项")
            return

        idx = sel[0]
        prize = self.prizes[idx]

        if prize.winners:
            if not messagebox.askyesno("确认", f"'{prize.name}' 已有中奖者，确定删除？\n中奖者将放回抽奖池"):
                return
            self.available.extend(prize.winners)

        del self.prizes[idx]
        if self.current_prize_idx >= len(self.prizes):
            self.current_prize_idx = max(0, len(self.prizes) - 1)

        self._update_prize_list()
        self._update_winners_display()
        self._update_participant_list()
        self.status_label.config(text=f"已删除奖项: {prize.name}")

    def _update_prize_list(self):
        self.prize_listbox.delete(0, tk.END)
        for i, p in enumerate(self.prizes):
            marker = "▶ " if i == self.current_prize_idx else "   "
            self.prize_listbox.insert(tk.END, f"{marker}{p.name} [{len(p.winners)}/{p.count}]")

        # 保持当前奖项选中状态
        if self.prizes and 0 <= self.current_prize_idx < len(self.prizes):
            self.prize_listbox.selection_set(self.current_prize_idx)

        self._update_current_prize_display()

    def _on_prize_select(self, event):
        sel = self.prize_listbox.curselection()
        if sel:
            self.current_prize_idx = sel[0]
            self._update_prize_list()
            self._clear_cards()

    def _prev_prize(self):
        if self.prizes and self.current_prize_idx > 0:
            self.current_prize_idx -= 1
            self._update_prize_list()
            self._clear_cards()
            self._update_fullscreen()

    def _next_prize(self):
        if self.prizes and self.current_prize_idx < len(self.prizes) - 1:
            self.current_prize_idx += 1
            self._update_prize_list()
            self._clear_cards()
            self._update_fullscreen()

    def _update_current_prize_display(self):
        self.current_prize_label.config(text=self._get_prize_text())
        self.remaining_label.config(text=self._get_remaining_text())
        self._update_fullscreen()

    # ==================== 抽奖逻辑 ====================
    def _prepare_draw_cards(self):
        """洗牌后准备抽奖卡片"""
        if not self._check_can_draw():
            return

        prize = self.prizes[self.current_prize_idx]
        if prize.remaining <= 0:
            self.hint_label.config(text=f"'{prize.name}' 已抽完，请切换奖项")
            self.status_label.config(text=f"'{prize.name}' 已抽完")
            return

        # 计算需要展示的卡片数量
        count = min(prize.remaining, len(self.available))

        self._clear_cards()

        # 预先分配中奖者（在 _clear_cards 之后，避免被清空）
        winners = random.sample(self.available, count)
        self.pending_winners = winners.copy()

        self._create_draw_cards(winners)
        self.cards_ready = True

        # 确保卡片容器居中显示
        if self.is_fullscreen and hasattr(self, 'fs_canvas'):
            self.fullscreen_window.update_idletasks()
            w, h = self.fs_canvas.winfo_width(), self.fs_canvas.winfo_height()
            self.fs_canvas.coords(self.fs_canvas_window, w // 2, h // 2)
        else:
            self.root.update_idletasks()
            w, h = self.display_canvas.winfo_width(), self.display_canvas.winfo_height()
            self.display_canvas.coords(self.canvas_window, w // 2, h // 2)

        self._update_hint(f"👆 共 {count} 张卡片，点击翻牌或点击「全部翻开」")
        self.status_label.config(text=f"洗牌完成！点击卡片抽取 {prize.name}")

    def _draw_one(self):
        """开始抽奖 - 如果没有卡片则先洗牌"""
        if self.shuffling:
            return

        # 如果没有准备好的卡片，先洗牌
        if not self.cards_ready or not self.flip_cards:
            if self._check_can_draw():
                self._shuffle_animation()
            return

        # 找到第一张未翻开的卡片并翻开
        for card in self.flip_cards:
            if not card.is_flipped:
                card.animate_flip()
                break

    def _get_active_window(self):
        """获取当前活动的窗口（用于 after 调用）"""
        if self.is_fullscreen and self.fullscreen_window:
            return self.fullscreen_window
        return self.root

    def _draw_all(self):
        """全部翻开"""
        if self.shuffling:
            return

        # 如果没有准备好的卡片，先洗牌
        if not self.cards_ready or not self.flip_cards:
            if self._check_can_draw():
                self._shuffle_animation()
            return

        # 翻开所有未翻开的卡片
        unflipped = [c for c in self.flip_cards if not c.is_flipped]
        if not unflipped:
            return

        # 依次快速翻开
        window = self._get_active_window()
        def flip_next(idx):
            if idx < len(unflipped):
                unflipped[idx].animate_flip()
                window.after(150, lambda: flip_next(idx + 1))

        flip_next(0)

    def _check_can_draw(self) -> bool:
        if not self.participants:
            messagebox.showwarning("提示", "请先导入名单")
            return False
        if not self.prizes:
            messagebox.showwarning("提示", "请先设置奖项")
            return False
        if not self.available:
            messagebox.showinfo("提示", "所有人都已中奖")
            return False
        return True

    def _update_hint(self, text: str):
        """更新提示标签（同时更新普通和全屏模式）"""
        self.hint_label.config(text=text)
        if self.is_fullscreen and hasattr(self, 'fs_hint_label'):
            self.fs_hint_label.config(text=text)

    def _on_card_flipped(self, card: FlipCard):
        """卡片翻开后的回调 - 确认中奖"""
        # 将该中奖者从pending移到正式中奖名单
        if card.name in self.pending_winners:
            self.pending_winners.remove(card.name)
            self.available.remove(card.name)
            prize = self.prizes[self.current_prize_idx]
            prize.winners.append(card.name)

            self._update_prize_list()
            self._update_winners_display()
            self._update_participant_list()

        # 检查是否所有卡片都翻开了
        all_flipped = all(c.is_flipped for c in self.flip_cards)
        if all_flipped:
            self._trigger_celebration()
            prize = self.prizes[self.current_prize_idx]
            names = ", ".join([c.name for c in self.flip_cards[:3]])
            if len(self.flip_cards) > 3:
                names += f" 等{len(self.flip_cards)}人"
            self.status_label.config(text=f"🎉 恭喜 {names} 获得 {prize.name}!")
            self._update_hint("🎊 本轮抽奖完成！可继续洗牌抽取或切换奖项")
            self.cards_ready = False
            self.pending_winners.clear()
        else:
            remaining = sum(1 for c in self.flip_cards if not c.is_flipped)
            flipped_count = len(self.flip_cards) - remaining
            self._update_hint(f"已翻开 {flipped_count} 张，还剩 {remaining} 张待翻开")

    def _clear_cards(self):
        # 清除卡片引用
        self.flip_cards.clear()
        self.cards_ready = False
        self.pending_winners.clear()

        # 清空普通模式卡片容器的所有子元素
        for child in self.cards_container.winfo_children():
            child.destroy()
        self.hint_label.config(text="")

        # 清空全屏模式卡片容器的所有子元素（安全检查）
        try:
            if hasattr(self, 'fs_cards_container') and self.fs_cards_container.winfo_exists():
                for child in self.fs_cards_container.winfo_children():
                    child.destroy()
                if hasattr(self, 'fs_hint_label') and self.fs_hint_label.winfo_exists():
                    self.fs_hint_label.config(text="")
        except tk.TclError:
            pass  # 全屏窗口已关闭，忽略错误

    def _get_active_cards_container(self):
        """获取当前活动的卡片容器"""
        if self.is_fullscreen and hasattr(self, 'fs_cards_container'):
            return self.fs_cards_container
        return self.cards_container

    def _calculate_card_size(self, count: int) -> Tuple[int, int, int]:
        """根据画布大小和卡片数量计算最佳卡片尺寸，返回 (宽, 高, 列数)"""
        canvas = self._get_active_canvas()
        canvas_w = canvas.winfo_width()
        canvas_h = canvas.winfo_height()

        if count == 1:
            # 单张卡片，占画布的合适比例
            card_w = min(400, int(canvas_w * 0.35))
            card_h = int(card_w * 0.6)
            return card_w, card_h, 1

        # 多张卡片，计算最佳布局
        # 根据卡片数量确定列数
        if count <= 3:
            cols = count
        elif count <= 6:
            cols = 3
        elif count <= 10:
            cols = 5
        else:
            cols = min(6, count)

        rows = math.ceil(count / cols)

        # 计算可用空间（留出边距）
        available_w = canvas_w * 0.85
        available_h = canvas_h * 0.7

        # 计算每张卡片的最大尺寸
        max_card_w = int(available_w / cols) - 12  # 减去间距
        max_card_h = int(available_h / rows) - 12

        # 保持宽高比 5:3
        if max_card_w * 0.6 > max_card_h:
            card_h = max_card_h
            card_w = int(card_h / 0.6)
        else:
            card_w = max_card_w
            card_h = int(card_w * 0.6)

        # 设置最小和最大限制
        card_w = max(120, min(300, card_w))
        card_h = max(75, min(180, card_h))

        return card_w, card_h, cols

    def _create_draw_cards(self, names: List[str]):
        """创建抽奖卡片"""
        count = len(names)
        if count == 0:
            return

        # 使用当前活动的容器
        active_container = self._get_active_cards_container()

        # 计算最佳卡片大小
        card_w, card_h, cols = self._calculate_card_size(count)

        if count == 1:
            # 单张大卡片
            card = FlipCard(active_container, width=card_w, height=card_h,
                           name=names[0], on_flip_complete=self._on_card_flipped)
            card.pack(pady=15)
            self.flip_cards.append(card)
        else:
            # 多张卡片网格
            container = tk.Frame(active_container, bg=Theme.BG_DARK)
            container.pack()

            for i, name in enumerate(names):
                card = FlipCard(container, width=card_w, height=card_h,
                               name=name, on_flip_complete=self._on_card_flipped)
                card.grid(row=i//cols, column=i%cols, padx=6, pady=6)
                self.flip_cards.append(card)

    def _trigger_celebration(self):
        if self.is_fullscreen and hasattr(self, 'fs_canvas'):
            w, h = self.fs_canvas.winfo_width(), self.fs_canvas.winfo_height()
            self.fs_particle_system.create_celebration(w, h)
            self.fs_particle_system.start()
        else:
            w, h = self.display_canvas.winfo_width(), self.display_canvas.winfo_height()
            self.particle_system.create_celebration(w, h)
            self.particle_system.start()

    def _reset_current_prize(self):
        if not self.prizes:
            return
        prize = self.prizes[self.current_prize_idx]
        if not prize.winners:
            messagebox.showinfo("提示", "当前奖项还没有中奖者")
            return

        if messagebox.askyesno("确认", f"重置 '{prize.name}'？\n中奖者将放回抽奖池"):
            self.available.extend(prize.winners)
            prize.winners.clear()
            self._clear_cards()
            self._update_prize_list()
            self._update_winners_display()
            self._update_participant_list()

            # 执行洗牌动画
            self._shuffle_animation()

    def _update_winners_display(self):
        self.winners_text.config(state=tk.NORMAL)
        self.winners_text.delete(1.0, tk.END)
        for prize in self.prizes:
            if prize.winners:
                self.winners_text.insert(tk.END, f"【{prize.name}】\n")
                for i, w in enumerate(prize.winners, 1):
                    self.winners_text.insert(tk.END, f"  {i}. {w}\n")
                self.winners_text.insert(tk.END, "\n")
        self.winners_text.config(state=tk.DISABLED)

    def _export_winners(self):
        all_winners = [(p.name, w) for p in self.prizes for w in p.winners]
        if not all_winners:
            messagebox.showinfo("提示", "还没有中奖者")
            return

        filepath = filedialog.asksaveasfilename(
            title="导出", defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("文本", "*.txt")]
        )
        if not filepath:
            return

        try:
            with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
                if filepath.endswith('.csv'):
                    writer = csv.writer(f)
                    writer.writerow(['奖项', '中奖者'])
                    writer.writerows(all_winners)
                else:
                    for p in self.prizes:
                        if p.winners:
                            f.write(f"【{p.name}】\n")
                            for i, w in enumerate(p.winners, 1):
                                f.write(f"  {i}. {w}\n")
                            f.write("\n")
            messagebox.showinfo("成功", f"已导出: {filepath}")
        except Exception as e:
            messagebox.showerror("错误", str(e))

    # ==================== 配置 ====================
    def _save_config(self):
        config = {
            'prizes': [{'name': p.name, 'count': p.count, 'winners': p.winners} for p in self.prizes],
            'participants': self.participants,
            'available': self.available,
            'current_prize_idx': self.current_prize_idx
        }
        try:
            path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lottery_config.json')
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(config, f, ensure_ascii=False, indent=2)
        except:
            pass

    def _load_config(self):
        try:
            path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lottery_config.json')
            if os.path.exists(path):
                with open(path, 'r', encoding='utf-8') as f:
                    config = json.load(f)
                self.participants = config.get('participants', [])
                self.available = config.get('available', [])
                self.current_prize_idx = config.get('current_prize_idx', 0)
                for p in config.get('prizes', []):
                    self.prizes.append(Prize(name=p['name'], count=p['count'],
                                            winners=p.get('winners', [])))
                self._update_participant_list()
                self._update_prize_list()
                self._update_winners_display()
        except:
            pass

    def run(self):
        self.root.protocol("WM_DELETE_WINDOW", self._on_closing)
        self.root.mainloop()

    def _on_closing(self):
        self._save_config()
        self.root.destroy()


def main():
    root = tk.Tk()
    app = LotteryApp(root)
    app.run()


if __name__ == "__main__":
    main()
